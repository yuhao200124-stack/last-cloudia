import json
import re
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK = Path('/workspace/scratch/9ea94a8a38d1/upload/被动技能_整理中_v6 (1).xlsx')
OLD_SOURCE = Path('/workspace/scratch/9ea94a8a38d1/old-site/app/data/source-skills.json')
OVERRIDES = Path('/workspace/sites/last-cloudia-skill-table/old-effect-overrides.json')
OUTPUT = Path('/workspace/sites/last-cloudia-skill-table/dist/data.js')


def text(value):
    if value is None:
        return ''
    return str(value).strip()


def sc_key(value):
    raw = text(value)
    try:
        number = float(raw)
        return str(int(number)) if number.is_integer() else str(number)
    except ValueError:
        return raw


def normalize(value):
    return re.sub(r'\s+', '', text(value)).replace('（', '(').replace('）', ')')


old_payload = json.loads(OLD_SOURCE.read_text(encoding='utf-8'))
old_skills = old_payload['skills']
effect_overrides = json.loads(OVERRIDES.read_text(encoding='utf-8'))

by_position = {}
by_name_sc = defaultdict(list)
by_name = defaultdict(list)
for item in old_skills:
    by_position[(item['sheet'], int(item['sourceRow']), item['sourceSide'])] = item
    by_name_sc[(normalize(item['name']), sc_key(item['sc']))].append(item)
    by_name[normalize(item['name'])].append(item)


def best_old_skill(sheet, row, side, name, sc, effect):
    positioned = by_position.get((sheet, row, side))
    if positioned and normalize(positioned['name']) == normalize(name):
        return positioned
    candidates = by_name_sc.get((normalize(name), sc_key(sc)), [])
    if not candidates:
        candidates = by_name.get(normalize(name), [])
    if not candidates:
        return None
    target = normalize(effect)
    return max(candidates, key=lambda item: (
        item['id'] in effect_overrides,
        SequenceMatcher(None, target, normalize(item.get('effect', ''))).ratio(),
    ))


def cell_url(cell):
    if not cell.hyperlink:
        return ''
    return text(cell.hyperlink.target or cell.hyperlink.location)


def enrich(sheet, row, side, name, sc, effect, sources, mark, url):
    old = best_old_skill(sheet, row, side, name, sc, effect)
    if old:
        effect = effect_overrides.get(old['id'], old.get('effect', effect))
        sources = old.get('sources') or sources
        sc = old.get('sc') if not sc_key(sc) else sc
        skill_id = old['id']
    else:
        skill_id = f'{sheet}:{side}:{row}'
    if isinstance(sources, str):
        sources = [part.strip() for part in sources.splitlines() if part.strip()]
    return {
        'id': skill_id,
        'name': text(name),
        'sc': sc_key(sc),
        'effect': text(effect),
        'sources': sources,
        'url': url,
        'mark': text(mark),
    }


wb = load_workbook(WORKBOOK, data_only=False)
sheet_order = [ws.title for ws in wb.worksheets if not ws.title.startswith('复制用的')]
result = {'generatedFrom': WORKBOOK.name, 'sheetOrder': sheet_order, 'sheets': {}}

all_ws = wb['全部技能']
all_rows = []
for row in range(2, all_ws.max_row + 1):
    name = all_ws.cell(row, 1).value
    if not text(name):
        continue
    all_rows.append(enrich(
        '全部技能', row, 'all', name, all_ws.cell(row, 2).value,
        all_ws.cell(row, 3).value, all_ws.cell(row, 4).value,
        all_ws.cell(row, 5).value, cell_url(all_ws.cell(row, 1)),
    ))
result['sheets']['全部技能'] = {'kind': 'all', 'rows': all_rows}

for sheet_name in sheet_order[1:]:
    ws = wb[sheet_name]
    lanes = []
    for side, columns in [('left', (1, 2, 3, 4, 5, 6)), ('right', (9, 10, 11, 12, 13, 14))]:
        type_col, name_col, sc_col, effect_col, source_col, mark_col = columns
        entries = []
        current_type = ''
        pending_separator = False
        for row in range(3, ws.max_row + 1):
            raw_type = text(ws.cell(row, type_col).value)
            if raw_type:
                current_type = raw_type
            name = ws.cell(row, name_col).value
            if not text(name):
                if entries and not any(text(ws.cell(row, col).value) for col in columns):
                    pending_separator = True
                continue
            if text(name).startswith(('http://', 'https://')) and not any(text(ws.cell(row, col).value) for col in (sc_col, effect_col, source_col)):
                continue
            if pending_separator:
                entries.append({'separator': True, 'sourceRow': row})
                pending_separator = False
            entry = enrich(
                sheet_name, row, side, name, ws.cell(row, sc_col).value,
                ws.cell(row, effect_col).value, ws.cell(row, source_col).value,
                ws.cell(row, mark_col).value, cell_url(ws.cell(row, name_col)),
            )
            entry['type'] = current_type
            entries.append(entry)
        lanes.append({'side': side, 'rows': entries})
    result['sheets'][sheet_name] = {'kind': 'split', 'lanes': lanes}

result['stats'] = {
    'allSkills': len(all_rows),
    'categoryEntries': sum(
        sum(not row.get('separator', False) for row in lane['rows'])
        for name, sheet in result['sheets'].items() if name != '全部技能'
        for lane in sheet['lanes']
    ),
    'separators': sum(
        sum(row.get('separator', False) for row in lane['rows'])
        for name, sheet in result['sheets'].items() if name != '全部技能'
        for lane in sheet['lanes']
    ),
    'descriptionOverrides': len(effect_overrides),
}

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
OUTPUT.write_text('window.SKILL_DATA = ' + json.dumps(result, ensure_ascii=False, separators=(',', ':')) + ';\n', encoding='utf-8')
print(json.dumps(result['stats'], ensure_ascii=False))
