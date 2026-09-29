"""The skill table's layout as an Excel workbook the user edits and sends back.

  python3 scripts/game-skill-excel.py export <out.xlsx>
  python3 scripts/game-skill-excel.py import <in.xlsx> [--dry-run]
  node scripts/build-game-skill-table.mjs            # then rebuild dist/game-skill-data.js

Every row is identified by its 游戏编号 (the game's PassiveSkillMst id); the calculator only ever uses that number.
- Taken from the workbook: the sheets and their order, the rows and their order, separators, 栏, 分组, 评价, and the
  名称 / 效果说明 texts (these change only the 简体 display; the 繁体 game text never changes). A 名称 or 效果说明
  left empty goes back to the game's own text.
- Always from the game data, whatever the workbook says: SC and 可学圣物 (changes there are only reported).
- 大类 / 标签 / 计算器: the row's classification, the same one the site keeps hidden on every row (scripts/build-skill-classes.mjs
  → docs/skill-classes-draft.json, by game number). View only: shown so the user can see them while moving rows; edits are
  only reported.
The import writes docs/game-skill-layout.json only when nothing is wrong; otherwise it lists the problems and stops.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
LAYOUT = ROOT / 'docs/game-skill-layout.json'
GAME = ROOT / 'docs/game-relic-passives.json'
SEPARATOR = '——分隔——'
INFO_SHEET = '说明'
CLASS = ['大类', '标签', '计算器']              # view only, from the game data by game number
COLS = {'all': ['游戏编号', '名称', 'SC', '效果说明', '可学圣物', '评价', *CLASS],
        'lanes': ['栏', '游戏编号', '名称', 'SC', '效果说明', '可学圣物', '评价', *CLASS],
        'split': ['栏', '分组', '游戏编号', '名称', 'SC', '效果说明', '可学圣物', '评价', *CLASS]}
WIDTH = {'栏': 8, '分组': 14, '游戏编号': 11, '名称': 22, 'SC': 6, '效果说明': 60, '可学圣物': 30, '评价': 7,
         '大类': 16, '标签': 28, '计算器': 14}
FIXED = {'游戏编号', 'SC', '可学圣物', *CLASS}    # never taken from the workbook
CLASSES = ROOT / 'docs/skill-classes-draft.json'
FONT = 'Arial'


def load():
    layout = json.loads(LAYOUT.read_text(encoding='utf-8'))
    game = {g['gameId']: g for g in json.loads(GAME.read_text(encoding='utf-8'))}
    if CLASSES.exists():
        for c in json.loads(CLASSES.read_text(encoding='utf-8'))['skills']:
            if c['id'] in game:
                game[c['id']]['cls'] = c
    return layout, game


def shown(layout, game, gid):
    own = layout.get('skills', {}).get(str(gid), {})
    g = game[gid]
    c = g.get('cls') or {}
    return {'名称': own.get('name') or g['nameS'], '效果说明': own.get('effect') or g['effectS'], '评价': own.get('mark', ''),
            'SC': g['sc'], '可学圣物': '\n'.join(f"{r['nameS']}（{r['rarity']}）" for r in g['relics']),
            '大类': '、'.join(c.get('cats', [])), '标签': '、'.join(c.get('tags', [])) or '—', '计算器': c.get('calc', '')}


# ---------------------------------------------------------------- export
def export(out):
    layout, game = load()
    wb = Workbook()
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_edit = PatternFill('solid', fgColor='FFF2CC')
    head_fixed = PatternFill('solid', fgColor='D9D9D9')
    cell_fixed = PatternFill('solid', fgColor='F2F2F2')
    sep_fill = PatternFill('solid', fgColor='DDEBF7')

    info = wb.active
    info.title = INFO_SHEET
    lines = [
        ('游戏数据技能表 · 排版', True),
        ('', False),
        ('每一行按“游戏编号”算：编号就是这个技能在游戏数据里的号码，计算器只看编号。换成别的编号 = 换成别的技能。', False),
        ('', False),
        ('可以改（表头黄色）', True),
        ('· 行的顺序、增加或删除行（增加时填游戏编号即可，名称等可以空着）', False),
        ('· 名称、效果说明：只改“简体”显示；“游戏原文（繁）”始终是游戏数据。清空这一格 = 恢复游戏原文。同一个技能在几个页里出现时，改一处就行。', False),
        ('· 评价（SS / S / A / B / C 或空）、分组、栏', False),
        ('· 工作表：可以增删、改名、调整顺序（这一页“说明”除外）。每个工作表就是网页上的一个分页。', False),
        ('', False),
        ('不会生效（表头灰色，永远按游戏数据）', True),
        ('· SC、可学圣物：改了也不会用，导入时会列出来告诉你', False),
        ('· 大类、标签、计算器：网站后台给每个技能记的分类，按游戏编号从游戏数据自动生成，只供查看。移动整行时它们跟着走；改了不会生效，导入时会列出来', False),
        ('· 标签“—”表示没有标签。计算器：能算 / 看条件（满足条件才吃得到）/ 不影响每段伤害 / 待确认', False),
        ('', False),
        ('移动技能', True),
        ('· 请整行移动（选中行号剪切、插入），不要只挪某一格：网站按这一行的游戏编号认技能', False),
        ('', False),
        ('栏和分组', True),
        ('· 分两栏的页：栏写 1 或 2（左栏 / 右栏）；分组写在“分组”列，相邻同名的行合成一组', False),
        ('· “基础属性”这类按项目分段的页：栏写段名（HP、MP、攻击力……），没有分组列', False),
        ('· “全部技能”这类单表页：没有栏和分组列', False),
        (f'· 分隔行：游戏编号那一格写“{SEPARATOR}”，其他格空着（分两栏的页要写栏）', False),
        ('', False),
        ('例子（分两栏的页）', True),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = info.cell(row=i, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=12 if i == 1 else 10)
    ex_row = len(lines) + 1
    example = [['栏', '分组', '游戏编号', '名称', 'SC', '效果说明', '可学圣物', '评价', *CLASS],
               [1, '暴伤', 700, '暴击伤害提升', '（自动）', '（自动显示游戏数据）', '（自动）', 'A', '（自动）', '（自动）', '（自动）'],
               [1, '', SEPARATOR, '', '', '', '', '', '', '', ''],
               [2, '暴击相关', 20800, '', '（自动）', '', '（自动）', '', '（自动）', '（自动）', '（自动）']]
    for r, row in enumerate(example):
        for col, v in enumerate(row, 1):
            c = info.cell(row=ex_row + r, column=col + 1, value=v)
            c.font = Font(name=FONT, bold=r == 0, size=10)
            c.border = border
            if r == 0:
                c.fill = head_fixed if v in FIXED else head_edit
    info.cell(row=ex_row + 4, column=2, value='第 4 行名称空着：显示游戏原文。').font = Font(name=FONT, size=10, color='808080')
    info.column_dimensions['A'].width = 4
    for col, w in zip('BCDEFGHIJKL', [8, 12, 11, 18, 8, 22, 10, 7, 10, 10, 10]):
        info.column_dimensions[col].width = w
    # the long text sits in column A; let it overflow to the right
    for i in range(1, len(lines) + 1):
        info.cell(row=i, column=1).alignment = Alignment(wrap_text=False)

    for name in layout['sheetOrder']:
        sh = layout['sheets'][name]
        kind = sh['kind']
        cols = COLS[kind]
        ws = wb.create_sheet(name)
        for col, h in enumerate(cols, 1):
            c = ws.cell(row=1, column=col, value=h)
            c.font = Font(name=FONT, bold=True, size=10)
            c.fill = head_fixed if h in FIXED else head_edit
            c.border = border
            c.alignment = Alignment(horizontal='center', vertical='center')
            ws.column_dimensions[c.column_letter].width = WIDTH[h]
        ws.freeze_panes = 'A2'
        lanes = [{'label': None, 'rows': sh['rows']}] if kind == 'all' else sh['lanes']
        r = 2
        for li, lane in enumerate(lanes):
            for row in lane['rows']:
                vals = {}
                if kind == 'lanes':
                    vals['栏'] = lane['label']
                elif kind == 'split':
                    vals['栏'] = li + 1
                if row.get('separator'):
                    vals['游戏编号'] = SEPARATOR
                else:
                    vals.update(shown(layout, game, row['id']))
                    vals['游戏编号'] = row['id']
                    if kind == 'split':
                        vals['分组'] = row.get('group', '')
                for col, h in enumerate(cols, 1):
                    c = ws.cell(row=r, column=col, value=vals.get(h, '') if vals.get(h, '') != '' else None)
                    c.font = Font(name=FONT, size=10, color='595959' if h in FIXED else '000000')
                    c.border = border
                    c.alignment = Alignment(vertical='center', wrap_text=h in ('效果说明', '可学圣物', '大类', '标签'),
                                            horizontal='center' if h in ('栏', 'SC', '评价', '游戏编号') else 'left')
                    if row.get('separator'):
                        c.fill = sep_fill
                    elif h in FIXED:
                        c.fill = cell_fixed
                r += 1
        ws.auto_filter.ref = f'A1:{ws.cell(row=1, column=len(cols)).column_letter}{r - 1}'
    wb.save(out)
    print(f'exported {len(layout["sheetOrder"])} sheets to {out}')


# ---------------------------------------------------------------- import
def text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace('\r\n', '\n').replace('\r', '\n').strip()


def read_workbook(path, game):
    wb = load_workbook(path, data_only=True)
    problems, notes = [], []
    order, sheets, seen = [], {}, {}      # seen[gid] -> list of (sheet, row, {名称, 效果说明, 评价, SC, 可学圣物})
    for ws in wb.worksheets:
        if ws.title == INFO_SHEET:
            continue
        header = [text(c.value) for c in ws[1]]
        pos = {h: i for i, h in enumerate(header) if h}
        if '游戏编号' not in pos:
            problems.append(f'「{ws.title}」第 1 行没有“游戏编号”这一列')
            continue
        rows = []
        for r, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            get = lambda h: text(cells[pos[h]]) if h in pos and pos[h] < len(cells) else ''
            raw = get('游戏编号')
            if not raw and not any(text(v) for v in cells):
                continue
            lane = get('栏')
            if raw.startswith('—') or raw.startswith('-') or raw == '分隔':
                rows.append({'r': r, 'sep': True, 'lane': lane})
                continue
            try:
                gid = int(raw)
            except ValueError:
                problems.append(f'「{ws.title}」第 {r} 行：游戏编号“{raw}”不是数字')
                continue
            if gid not in game:
                problems.append(f'「{ws.title}」第 {r} 行：游戏数据里没有编号 {gid}（{get("名称") or "无名称"}）')
                continue
            rows.append({'r': r, 'id': gid, 'lane': lane, 'group': get('分组')})
            seen.setdefault(gid, []).append((ws.title, r, {h: get(h) for h in ('名称', '效果说明', '评价', 'SC', '可学圣物', *CLASS)}))
        lanes_used = [x['lane'] for x in rows if x['lane']]
        groups_used = [x for x in rows if x.get('group')]
        if not lanes_used and not groups_used:
            kind = 'all'
        elif lanes_used and all(v in ('1', '2') for v in lanes_used):
            kind = 'split'
        elif lanes_used and not any(v in ('1', '2') for v in lanes_used):
            kind = 'lanes'
            if groups_used:
                notes.append(f'「{ws.title}」按项目分段，没有分组：“分组”列不会显示')
        else:
            problems.append(f'「{ws.title}」的“栏”混用了 1/2 和段名，请统一（分两栏写 1 或 2；分段写段名）')
            continue
        if kind == 'all':
            sheet = {'kind': 'all', 'rows': [{'separator': True} if x.get('sep') else {'id': x['id']} for x in rows]}
        elif kind == 'split':
            lanes = [{'label': None, 'rows': []}, {'label': None, 'rows': []}]
            for x in rows:
                if not x['lane']:
                    problems.append(f'「{ws.title}」第 {x["r"]} 行没有写栏（1 或 2）')
                    continue
                lanes[int(x['lane']) - 1]['rows'].append({'separator': True} if x.get('sep') else {'id': x['id'], **({'group': x['group']} if x['group'] else {})})
            sheet = {'kind': 'split', 'lanes': lanes}
        else:
            labels = []
            for x in rows:
                if not x['lane']:
                    problems.append(f'「{ws.title}」第 {x["r"]} 行没有写栏（段名）')
                elif x['lane'] not in labels:
                    labels.append(x['lane'])
            sheet = {'kind': 'lanes', 'lanes': [{'label': l, 'rows': [{'separator': True} if x.get('sep') else {'id': x['id']} for x in rows if x['lane'] == l]} for l in labels]}
        order.append(ws.title)
        sheets[ws.title] = sheet
    return order, sheets, seen, problems, notes


def importing(path, dry):
    layout, game = load()
    order, sheets, seen, problems, notes = read_workbook(path, game)
    skills, changes = {}, {'名称': [], '效果说明': [], '评价': []}
    field_key = {'名称': 'name', '效果说明': 'effect', '评价': 'mark'}
    game_text = {'名称': 'nameS', '效果说明': 'effectS'}
    for gid in sorted(game):
        before = shown(layout, game, gid)
        own = dict(layout.get('skills', {}).get(str(gid), {}))
        for f in ('名称', '效果说明', '评价'):
            vals = [(s, r, v[f]) for s, r, v in seen.get(gid, [])]
            new = {v for _, _, v in vals if v != before[f]}
            if len(new) > 1:
                where = '；'.join(f'「{s}」第{r}行：{v or "（空）"}' for s, r, v in vals if v != before[f])
                problems.append(f'编号 {gid}（{before["名称"]}）的{f}在几处改得不一样：{where}')
                continue
            if not new:
                continue
            value = new.pop()
            if f == '评价':
                own['mark'] = value
            elif not value or value == game[gid][game_text[f]]:
                if field_key[f] not in own:
                    continue              # already the game text
                own.pop(field_key[f])
                value = value or '（恢复游戏原文）'
            else:
                own[field_key[f]] = value
            changes[f].append(f'{gid} {before["名称"]}：{before[f] or "（空）"} → {value or "（空）"}')
        for f in ('SC', '可学圣物', *CLASS):
            for s, r, v in seen.get(gid, []):
                if v[f] and v[f] != str(before[f]):
                    notes.append(f'「{s}」第 {r} 行：{f}改成了“{v[f]}”，不会生效（按游戏数据：{before[f] or "（空）"}）')
        own = {k: v for k, v in own.items() if v}
        if own:
            skills[str(gid)] = own
    missing = [gid for gid in game if gid not in seen]
    if missing:
        notes.append(f'有 {len(missing)} 个游戏被动没放在任何工作表里，会自动放到“全部技能”末尾和“杂项”的“新增待排”：' + '、'.join(f'{g} {game[g]["nameS"]}' for g in missing[:30]) + ('……' if len(missing) > 30 else ''))
    report = {'sheets': order, 'problems': problems, 'notes': notes, 'changes': {k: v for k, v in changes.items() if v}}
    if not problems and not dry:
        out = {'note': layout.get('note', ''), 'sheetOrder': order, 'sheets': sheets, 'skills': skills}
        LAYOUT.write_text(json.dumps(out, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
        report['written'] = str(LAYOUT.relative_to(ROOT))
    print(json.dumps(report, ensure_ascii=False, indent=1))
    return 1 if problems else 0


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == 'export':
        export(sys.argv[2])
    elif len(sys.argv) >= 3 and sys.argv[1] == 'import':
        sys.exit(importing(sys.argv[2], '--dry-run' in sys.argv))
    else:
        print(__doc__)
        sys.exit(2)
