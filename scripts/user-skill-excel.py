"""The user's own skill-table workbook as the home page's layout (user 2026-09-29: “之后我要调整格式或者位置我就改这个改完发你你再发布”).

  python3 scripts/user-skill-excel.py import <file.xlsx> --dry-run   # check, list what changes on the site
  python3 scripts/user-skill-excel.py import <file.xlsx>             # write docs/game-skill-layout.json
  node scripts/build-game-skill-table.mjs                            # then rebuild dist/game-skill-data.js

The workbook is the user's format (the table they have always kept), plus the game number:
- 全部技能: 技能名称 / SC / 技能效果 / 可学习圣物（中文）/ 评价 / 游戏编号 / 游戏里的名字（简）, headers in row 1.
- every other sheet: two lanes, columns A–H and I–P (技能类型, 技能名称, SC, 技能效果/说明, 可学习圣物, 评价, 游戏编号,
  游戏里的名字); row 1 is the title, row 2 the header (a sheet without the header row starts at row 2).
  技能类型 carries down to the rows below until the next one (the group); a blank row between skills is a separator.
- sheets named 说明, 缺少的技能 (game skills the table does not have yet, for the user to place) or 复制用的… are skipped.
Every row is identified by its 游戏编号 (a row with a number and no name counts: the name comes from the game data and
the row is reported). A row without a number is looked up by its 技能名称 in docs/user-skill-names.json
(and reported); a row that cannot be found is a problem, with the likeliest skills listed. Taken from the workbook:
the sheets and their order, the lanes, row order, groups, separators, and 评价 (from 全部技能 only — the other sheets
use that column for notes). Names, SC, effects and relics always come from the game data. A sheet the site has but the
workbook does not (基础属性) is kept as it is. The import writes only when nothing is wrong; it then adds the workbook's
names to docs/user-skill-names.json.
"""
import collections
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
LAYOUT = ROOT / 'docs/game-skill-layout.json'
GAME = ROOT / 'docs/game-relic-passives.json'
NAMES = ROOT / 'docs/user-skill-names.json'
RELICS = ROOT / 'docs/user-skill-relics.json'
SKIP = ('说明', '缺少的技能')


def text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace('\r\n', '\n').replace('\r', '\n').strip()


# ---------------------------------------------------------------- likely skills for a row without a number
ROMAN = {'2': 'II', '3': 'III', '4': 'IV', '5': 'V', '6': 'VI'}


def norm(s):
    s = re.sub(r'[\s　]', '', s).replace('（', '(').replace('）', ')').replace('【', '[').replace('】', ']').replace('・', '·')
    return re.sub(r'(\D)([2-6])$', lambda m: m.group(1) + ROMAN[m.group(2)], s).lower()


def dice(a, b):
    grams = lambda s: collections.Counter(s[i:i + 2] for i in range(len(s) - 1)) if len(s) > 1 else collections.Counter([s])
    A, B = grams(a), grams(b)
    total = sum(A.values()) + sum(B.values())
    return 2 * sum((A & B).values()) / total if total else 0


def clean(s):
    return re.sub(r'[\s，。、,.：:；;!！（）()\[\]【】「」『』“”"\'・･·+\-%％]', '', s)


def likely(row, game, relics):
    """the three game skills most like a row, by name, relics, SC and effect text"""
    rel = [re.sub(r'[（(][^）)]*[）)]\s*$', '', p).strip() for p in re.split(r'[\n、]', row['relics']) if p.strip()]
    rel = {relics.get(r, r) for r in rel}
    scored = []
    for gid, g in game.items():
        grel = {x['nameS'] for x in g['relics']}
        s = (dice(norm(row['name']), norm(g['nameS'])) + 1.5 * (len(rel & grel) / max(len(rel), len(grel), 1))
             + (1 if str(g['sc']) == row['sc'] else 0) + (dice(clean(row['effect']), clean(g['effectS'])) if row['effect'] else 0))
        scored.append((s, gid))
    return [gid for _, gid in sorted(scored, reverse=True)[:3]]


# ---------------------------------------------------------------- read
def read_workbook(path, game, names):
    wb = load_workbook(path, data_only=True)
    order, sheets, rows, problems, notes = [], {}, [], [], []

    def resolve(sheet, r, cells):
        raw, name = cells['id'], cells['name']
        where = f'「{sheet}」第 {r} 行' + (f'「{name}」' if name else '')
        if raw:
            try:
                gid = int(raw)
            except ValueError:
                problems.append(f'{where}：游戏编号“{raw}”不是数字')
                return None
            if gid not in game:
                problems.append(f'{where}：游戏数据里没有编号 {gid}')
                return None
            if not name:
                notes.append(f'{where}：技能名称是空的，按编号 {gid} 补上名字「{game[gid]["nameS"]}」')
                return gid
            known = names.get(name)
            if known is not None and known != gid:
                notes.append(f'{where}：编号是 {gid}（{game[gid]["nameS"]}），但这个名字以前对应 {known}（{game[known]["nameS"]}）。是不是只挪了某一格？')
            return gid
        if name in names:
            gid = names[name]
            notes.append(f'{where}：没有游戏编号，按名字补成 {gid}（{game[gid]["nameS"]}）')
            return gid
        problems.append(f'{where}：没有游戏编号，名字也对不上。最像的是：' + '、'.join(f'{g} {game[g]["nameS"]}' for g in likely(cells, game, RELICS_MAP)))
        return None

    for ws in wb.worksheets:
        if ws.title in SKIP or ws.title.startswith('复制用的'):
            continue
        if ws.title == '全部技能':
            head = {text(c.value): c.column for c in ws[1] if text(c.value)}
            col = {'name': head.get('技能名称', 1), 'sc': head.get('SC', 2), 'effect': head.get('技能效果', 3),
                   'relics': next((v for k, v in head.items() if k.startswith('可学习圣物')), 4), 'mark': head.get('评价', 5),
                   'id': head.get('游戏编号')}
            if not col['id']:
                problems.append('「全部技能」第 1 行没有“游戏编号”这一列')
                continue
            out = []
            for r in range(2, ws.max_row + 1):
                cells = {k: text(ws.cell(r, c).value) for k, c in col.items()}
                # a row is a skill when it has a name or a game number (the number decides; user 2026-09-30)
                if not cells['name'] and not cells['id']:
                    continue
                gid = resolve(ws.title, r, cells)
                if gid is not None:
                    out.append({'id': gid})
                    rows.append((ws.title, r, gid, cells))
            sheets[ws.title] = {'kind': 'all', 'rows': out}
        else:
            first = 3 if text(ws.cell(2, 2).value) == '技能名称' else 2
            lanes = []
            for start in (1, 9):
                keys = ('type', 'name', 'sc', 'effect', 'relics', 'mark', 'id')
                out, group, pending = [], '', False
                for r in range(first, ws.max_row + 1):
                    cells = {k: text(ws.cell(r, start + i).value) for i, k in enumerate(keys)}
                    if cells['type']:
                        group = cells['type']
                    if not cells['name'] and not cells['id']:
                        if out and not any(cells.values()):
                            pending = True
                        continue
                    if pending:
                        out.append({'separator': True})
                        pending = False
                    gid = resolve(ws.title, r, cells)
                    if gid is not None:
                        out.append({'id': gid, **({'group': group} if group else {})})
                        rows.append((ws.title, r, gid, cells))
                lanes.append({'label': None, 'rows': out})
            if not any(l['rows'] for l in lanes):
                notes.append(f'「{ws.title}」是空的，跳过')
                continue
            sheets[ws.title] = {'kind': 'split', 'lanes': lanes}
        order.append(ws.title)
    return order, sheets, rows, problems, notes


RELICS_MAP = {}


def ids_of(sheet):
    rows = sheet['rows'] if sheet['kind'] == 'all' else [x for l in sheet['lanes'] for x in l['rows']]
    return [x['id'] for x in rows if not x.get('separator')]


def importing(path, dry):
    global RELICS_MAP
    layout = json.loads(LAYOUT.read_text(encoding='utf-8'))
    game = {g['gameId']: g for g in json.loads(GAME.read_text(encoding='utf-8'))}
    names = json.loads(NAMES.read_text(encoding='utf-8'))['names'] if NAMES.exists() else {}
    RELICS_MAP = json.loads(RELICS.read_text(encoding='utf-8'))['relics'] if RELICS.exists() else {}
    order, sheets, rows, problems, notes = read_workbook(path, game, names)
    if '全部技能' not in sheets:
        problems.append('没有“全部技能”工作表')
    else:
        dup = [g for g, n in collections.Counter(ids_of(sheets['全部技能'])).items() if n > 1]
        for g in dup:
            where = '、'.join(f'第 {r} 行「{c["name"]}」' for s, r, gid, c in rows if s == '全部技能' and gid == g)
            problems.append(f'「全部技能」里编号 {g}（{game[g]["nameS"]}）出现了几次：{where}')
    # sheets the site has and the workbook does not (基础属性) stay where they are
    for i, name in enumerate(layout['sheetOrder']):
        if name not in sheets:
            before = next((n for n in reversed(layout['sheetOrder'][:i]) if n in order), None)
            order.insert(order.index(before) + 1 if before else 0, name)
            sheets[name] = layout['sheets'][name]
            notes.append(f'「{name}」不在表里，网站上保留现在的样子')
    placed = {g for s in sheets.values() for g in ids_of(s)}
    missing = [g for g in game if g not in placed]
    if missing:
        notes.append(f'有 {len(missing)} 个游戏被动没放在任何工作表里，会自动放到“全部技能”最后和“杂项”的“新增待排”：' + '、'.join(f'{g} {game[g]["nameS"]}' for g in missing))
    # 评价 from 全部技能; names / effects rewritten earlier through the old workbook (game-skill-excel.py, deleted 2026-09-30; none today) stay
    marks = {gid: c['mark'] for s, r, gid, c in rows if s == '全部技能' and c['mark']}
    skills, mark_changes = {}, []
    for gid in sorted(game):
        own = {k: v for k, v in layout.get('skills', {}).get(str(gid), {}).items() if k != 'mark'}
        before, after = layout.get('skills', {}).get(str(gid), {}).get('mark', ''), marks.get(gid, '')
        if before != after:
            mark_changes.append(f'{gid} {game[gid]["nameS"]}：{before or "（空）"} → {after or "（空）"}')
        if after:
            own['mark'] = after
        if own:
            skills[str(gid)] = own
    # what changes on the site, per sheet
    changes = []
    for name in order:
        new, old = sheets[name], layout['sheets'].get(name)
        if old is None:
            changes.append(f'「{name}」新分页，{len(ids_of(new))} 个技能')
            continue
        if new == old:
            continue
        a, b = collections.Counter(ids_of(new)), collections.Counter(ids_of(old))
        plus, minus = sum((a - b).values()), sum((b - a).values())
        parts = [f'多了 {plus} 个' if plus else '', f'少了 {minus} 个' if minus else '']
        if not plus and not minus:
            parts.append('顺序、分组或分隔有变化')
        changes.append(f'「{name}」' + '，'.join(p for p in parts if p))
    for name in layout['sheetOrder']:
        if name not in order:
            changes.append(f'「{name}」删掉了')
    if [n for n in order if n in layout['sheetOrder']] != [n for n in layout['sheetOrder'] if n in order]:
        changes.append('分页顺序有变化')
    report = {'sheets': order, 'problems': problems, 'notes': notes, 'changes': changes, 'ratings': mark_changes}
    if not problems and not dry:
        out = {'note': layout.get('note', ''), 'sheetOrder': order, 'sheets': sheets, 'skills': skills}
        LAYOUT.write_text(json.dumps(out, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
        # remember the workbook's names for rows added later without a number
        for s, r, gid, c in rows:
            if c['name']:
                names.setdefault(c['name'], gid)
        NAMES.write_text(json.dumps({'note': json.loads(NAMES.read_text(encoding='utf-8'))['note'] if NAMES.exists() else '', 'names': dict(sorted(names.items()))}, ensure_ascii=False, indent=0) + '\n', encoding='utf-8')
        report['written'] = str(LAYOUT.relative_to(ROOT))
    print(json.dumps(report, ensure_ascii=False, indent=1))
    return 1 if problems else 0


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == 'import':
        sys.exit(importing(sys.argv[2], '--dry-run' in sys.argv))
    print(__doc__)
    sys.exit(2)
